"""Normalize the user's provider exports without altering the source workbooks.

Run through xlsx_repl:
  exec(compile(open(script).read(), script, "exec"))
  extract_benchmark_workbooks(source_directory, output_directory, "2026-08-31")

This intentionally supports these specific export layouts, not arbitrary Excel.
The separate release importer validates the package again before saving it.
"""

import calendar
import hashlib
import json
import math
from datetime import date, datetime
from pathlib import Path

import openpyxl


def extract_benchmark_workbooks(source_directory, output_directory, as_of):
    source = Path(source_directory)
    output = Path(output_directory)
    output.mkdir(parents=True, exist_ok=True)
    end = date.fromisoformat(as_of)
    if end.day != calendar.monthrange(end.year, end.month)[1]:
        raise ValueError("Reporting date must be a calendar month-end")
    files = ["Benchmark-fonddashboard_ex-NBP.xlsx", "Benchmark-pris_NBP.xlsx"]
    audit = {
        "asOf": as_of,
        "sources": [
            {"file": name, "sha256": hashlib.sha256((source / name).read_bytes()).hexdigest()}
            for name in files
        ],
        "currencyTreatment": "All supplied levels are already NOK; no FX conversion.",
        "nbpLevelColumn": "F: Total Return (Gross, Unhedged)",
        "nbpUnusedColumn": "G: Cumulative Return % (Gross, Unhedged)",
        "excluded": [],
        "series": [],
    }
    package = {"version": 1, "series": [], "assignments": {
        "F0GBR04NJK": "OSEFX_TR_NOK",
        "F00001GW7Y": "MSCI_ACWI_NET_TR_NOK",
        "F0000125XE": "MSCI_ACWI_NET_TR_NOK",
        "F00000ZFGS": "MSCI_ACWI_NET_TR_NOK",
        "F00001GU8B": "MSCI_EM_NET_TR_NOK",
        "F00000OAR8": "MSCI_EM_NET_TR_NOK",
        "F0GBR05THA": "MSCI_WORLD_NET_TR_NOK",
        "F00001EWFB": "BLOOMBERG_GLOBAL_AGG_TR_NOK_HEDGED",
        "F000014RM8": "BLOOMBERG_GLOBAL_AGG_TR_NOK_HEDGED",
        "F00000Z0Y3": "BLOOMBERG_GLOBAL_AGG_TR_NOK_HEDGED",
        "F00000LKIE": "NBP_NOHYNH",
        "F00001SEG2": "NBP_NORM123D3",
        "F00000QLFG": "NBP_NORMFRN",
    }}

    def collect(identifier, name, provider, return_type, hedging, filename,
                sheet, original_name, observations):
        levels, cells = [], []
        for date_cell, level_cell, cumulative_cell in observations:
            stamp, value = date_cell.value, level_cell.value
            if not isinstance(stamp, (date, datetime)):
                raise ValueError(f"{sheet.title}!{date_cell.coordinate}: not a date")
            day = stamp.date() if isinstance(stamp, datetime) else stamp
            if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value <= 0:
                raise ValueError(f"{sheet.title}!{level_cell.coordinate}: invalid index level")
            record = {
                "date": day.isoformat(), "level": value,
                "dateCell": date_cell.coordinate, "levelCell": level_cell.coordinate,
            }
            if day > end or day.day != calendar.monthrange(day.year, day.month)[1]:
                audit["excluded"].append({
                    "file": filename, "sheet": sheet.title, **record,
                    "reason": "After reporting date and/or not a complete calendar month-end",
                })
                continue
            levels.append({"date": day.isoformat(), "level": value})
            if cumulative_cell is not None:
                record["unusedCumulativeReturnCell"] = cumulative_cell.coordinate
                record["unusedCumulativeReturnPct"] = cumulative_cell.value
            cells.append(record)
        levels.sort(key=lambda p: p["date"])
        cells.sort(key=lambda p: p["date"])
        if not levels or levels[-1]["date"] != as_of:
            raise ValueError(f"{identifier}: missing reporting month-end")
        for previous, current in zip(levels, levels[1:]):
            a, b = date.fromisoformat(previous["date"]), date.fromisoformat(current["date"])
            if (b.year * 12 + b.month) - (a.year * 12 + a.month) != 1:
                raise ValueError(f"{identifier}: duplicate or missing month")
        package["series"].append({
            "id": identifier, "name": name, "provider": provider,
            "currency": "NOK", "returnType": return_type, "hedging": hedging,
            "levels": levels,
        })
        by_date = {p["date"]: p["level"] for p in levels}
        checks = {}
        offsets = {"1M": 1, "3M": 3, "6M": 6, "YTD": end.month,
                   "1Y": 12, "3Y": 36, "5Y": 60}
        for period, months in offsets.items():
            year, month0 = divmod(end.year * 12 + end.month - 1 - months, 12)
            month = month0 + 1
            start = date(year, month, calendar.monthrange(year, month)[1]).isoformat()
            if start in by_date:
                checks[period] = {
                    "start": start, "end": as_of, "startLevel": by_date[start],
                    "endLevel": by_date[as_of],
                    "returnPct": (by_date[as_of] / by_date[start] - 1) * 100,
                }
        audit["series"].append({
            "id": identifier, "file": filename, "sheet": sheet.title,
            "originalName": original_name,
            "start": levels[0]["date"], "end": levels[-1]["date"],
            "count": len(levels), "checks": checks, "observations": cells,
        })

    wb = openpyxl.load_workbook(source / files[0], data_only=True)
    ws = wb["Price"]
    if ws["A6"].value != "Currency: Norwegian Krone" or ws["A4"].value != "Frequency: Monthly":
        raise ValueError("Expected an explicitly NOK-denominated monthly export")
    specs = [
        (10, "Bloomberg Global Aggregate TR Hdg NOK (Market Price, NOK)",
         "BLOOMBERG_GLOBAL_AGG_TR_NOK_HEDGED", "Bloomberg Global Aggregate TR Hedged NOK", "Bloomberg", "TR", "NOK_HEDGED"),
        (11, "MSCI ACWI NR USD (Market Price, NOK)",
         "MSCI_ACWI_NET_TR_NOK", "MSCI ACWI Net TR NOK", "MSCI", "NET_TR", "UNHEDGED"),
        (12, "MSCI EM NR USD (Market Price, NOK)",
         "MSCI_EM_NET_TR_NOK", "MSCI EM Net TR NOK", "MSCI", "NET_TR", "UNHEDGED"),
        (13, "MSCI World NR USD (Market Price, NOK)",
         "MSCI_WORLD_NET_TR_NOK", "MSCI World Net TR NOK", "MSCI", "NET_TR", "UNHEDGED"),
        (14, "OSE Oslo Børs Mutual Fund TR NOK (Market Price, NOK)",
         "OSEFX_TR_NOK", "OSEFX Total Return NOK", "Euronext", "TR", "NA"),
    ]
    for row, original, identifier, name, provider, kind, hedge in specs:
        if ws.cell(row, 1).value != original:
            raise ValueError(f"Unexpected index name in Price!A{row}")
        collect(identifier, name, provider, kind, hedge, files[0], ws, original,
                [(ws.cell(9, col), ws.cell(row, col), None)
                 for col in range(2, ws.max_column + 1)])
    wb.close()

    wb = openpyxl.load_workbook(source / files[1], data_only=True)
    for code in ["NOHYNH", "NORM123D3", "NORMFRN"]:
        ws = wb[code]
        if not str(ws["A1"].value).startswith(f"Price History: {code}"):
            raise ValueError(f"Wrong index identity for {code}")
        if ws["A3"].value != "Date" or ws["F3"].value != "Total Return (Gross, Unhedged)":
            raise ValueError(f"Unexpected date/level columns for {code}")
        if ws["G3"].value != "Cumulative Return % (Gross, Unhedged)":
            raise ValueError(f"Unexpected cumulative return column for {code}")
        collect(f"NBP_{code}", f"NBP {code}", "Nordic Bond Pricing",
                "GROSS_TR", "UNHEDGED", files[1], ws, ws["A1"].value,
                [(ws.cell(row, 1), ws.cell(row, 6), ws.cell(row, 7))
                 for row in range(4, ws.max_row + 1)])
    wb.close()

    for filename, data in [("package.json", package), ("source-audit.json", audit)]:
        (output / filename).write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({
        "output": str(output), "asOf": as_of, "series": len(package["series"]),
        "assignments": len(package["assignments"]),
        "observations": sum(len(s["levels"]) for s in package["series"]),
        "excluded": audit["excluded"],
        "coverage": [{k: s[k] for k in ["id", "start", "end", "count"]} for s in audit["series"]],
        "checks": [{"id": s["id"], "returns": {p: v["returnPct"] for p, v in s["checks"].items()}}
                   for s in audit["series"]],
    }, ensure_ascii=False, indent=2))
    return package, audit
